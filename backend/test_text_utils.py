import tempfile
import os
from text_utils import extract_text, generate_quiz, generate_puzzles

def test_extract_text_txt():
    with tempfile.NamedTemporaryFile(mode="w+", suffix=".txt", delete=False) as f:
        f.write("Hello, pytest!")
        f.seek(0)
        path = f.name
    result = extract_text(path)
    assert result == "Hello, pytest!"
    os.remove(path)

def test_extract_text_empty():
    with tempfile.NamedTemporaryFile(mode="w+", suffix=".txt", delete=False) as f:
        path = f.name
    result = extract_text(path)
    assert result == ""
    os.remove(path)

def test_extract_text_docx():
    from docx import Document
    path = os.path.join(tempfile.gettempdir(), "test.docx")
    doc = Document()
    doc.add_paragraph("Docx test paragraph.")
    doc.save(path)
    result = extract_text(path)
    assert "Docx test paragraph." in result
    os.remove(path)

def test_generate_quiz_basic():
    text = "This is a test sentence with some nouns like apple and banana. Another sentence with cat and dog."
    quiz = generate_quiz(text)
    assert isinstance(quiz, list)
    assert len(quiz) > 0
    for q in quiz:
        assert "question" in q
        assert "choices" in q
        assert "answer" in q
        assert q["answer"] in q["choices"]

def test_generate_quiz_no_nouns():
    text = "Quickly running, always jumping, happily laughing."
    quiz = generate_quiz(text)
    assert quiz == []

def test_generate_quiz_short_sentences():
    text = "Cat. Dog. Apple."
    quiz = generate_quiz(text)
    assert quiz == []

def test_generate_puzzles_basic():
    text = "Python programming language provides powerful features."
    puzzles = generate_puzzles(text)
    assert isinstance(puzzles, list)
    assert len(puzzles) > 0
    for p in puzzles:
        assert "puzzle" in p
        assert "answer" in p
        scrambled = p["puzzle"].replace("Unscramble this word: ", "").replace(" ", "")
        assert sorted(scrambled.lower()) == sorted(p["answer"].lower())

def test_generate_puzzles_min_word_length():
    text = "Short small tiniest apple programming"
    puzzles = generate_puzzles(text)
    assert any(p["answer"] == "programming" for p in puzzles) or any(p["answer"] == "apple" for p in puzzles)

def test_extract_text_pdf():
    # Create a minimal PDF for test
    from fpdf import FPDF
    path = os.path.join(tempfile.gettempdir(), "test.pdf")
    pdf = FPDF()
    pdf.add_page()
    pdf.set_font("Arial", size=12)
    pdf.cell(200, 10, txt="Test PDF content.", ln=True)
    pdf.output(path)
    result = extract_text(path)
    assert "Test PDF content." in result
    os.remove(path)

def test_extract_text_malformed_docx():
    # Test extract_text handles corrupt/malformed docx gracefully - just passes without error
    path = os.path.join(tempfile.gettempdir(), "corrupt.docx")
    with open(path, "w") as f:
        f.write("not a real docx file")
    result = extract_text(path)
    assert result == ""
    os.remove(path)
